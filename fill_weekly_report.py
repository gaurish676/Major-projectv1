import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import copy
import subprocess
import os

TEMPLATE_PATH = 'weekly-progress-report_2026.docx.bak'
TARGET_DOCX = 'weekly-progress-report_2026.docx'
TARGET_PDF = 'weekly-progress-report_2026.pdf'

doc = docx.Document(TEMPLATE_PATH)

# 1. Clean up excessive blank paragraphs from original template
p_to_remove = []
for p in doc.paragraphs:
    if not p.text.strip():
        p_to_remove.append(p)

for p in p_to_remove:
    p._element.getparent().remove(p._element)

# 2. Update Header Information (Title, Date, Guide line, Signature)
for p in doc.paragraphs:
    if 'Project Title:' in p.text:
        p.text = 'Project Title: Creditz – Institutional Academic & Extra-Curricular Activity Credit Management System with AI-Assisted Verification and GraphRAG'
        p.runs[0].font.name = 'Times New Roman'
        p.runs[0].font.size = Pt(11)
        p.runs[0].font.bold = True
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
    elif 'Name of the Guide:' in p.text:
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(6)
    elif 'Date:' in p.text:
        p.text = 'Date: 28/09/2026'
        p.runs[0].font.name = 'Times New Roman'
        p.runs[0].font.size = Pt(10.5)
        p.paragraph_format.first_line_indent = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
    elif 'Signature of guide' in p.text:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(4)

t1 = doc.tables[0]

# 3. Create Table 2 by cloning Table 1 structure to ensure clean 2-page division (Weeks 1-2 on Page 1, Weeks 3-4 on Page 2)
t2_element = copy.deepcopy(t1._tbl)
t1._tbl.addnext(t2_element)

p_break = doc.add_paragraph()
p_break.paragraph_format.page_break_before = True
t1._tbl.addnext(p_break._element)

t2 = docx.table.Table(t2_element, doc)

# In t1: remove rows 9 to 16 (Weeks 3 & 4)
for r_idx in range(16, 8, -1):
    tr = t1.rows[r_idx]._tr
    tr.getparent().remove(tr)

# In t2: remove rows 1 to 8 (Weeks 1 & 2)
for r_idx in range(8, 0, -1):
    tr = t2.rows[r_idx]._tr
    tr.getparent().remove(tr)

# 4. Weekly Data Definitions
weeks_data = [
    {
        "table": t1,
        "start_row": 1,
        "week_str": "Week 1\nfrom:\n31/08/2026\nto:\n05/09/2026",
        "tasks": (
            "• Literature survey on AICTE/VTU 100/200 activity credit guidelines and degree clearance.\n"
            "• Formulated problem statement on manual verification delays, fraud, and cap violations.\n"
            "• Defined 3-tier user role hierarchy (Student, Faculty Mentor, Head of Department).\n"
            "• Designed SQLite schema for users, activity categories, submissions, and audit logs.\n"
            "• Initialized repo scaffolding with Vite, React 19, TypeScript, and Express REST server."
        ),
        "suggestions": (
            "• Formalize mathematical constraints for semester limits (30 pts ceiling) and category domain caps.\n"
            "• Ensure clear separation between authoritative database records and derived intelligence layers.\n"
            "• Benchmark UI requirements for non-technical faculty evaluators."
        ),
    },
    {
        "table": t1,
        "start_row": 5,
        "week_str": "Week 2\nfrom:\n07/09/2026\nto:\n12/09/2026",
        "tasks": (
            "• Implemented deterministic points engine enforcing 30 pts/sem limit and category caps.\n"
            "• Built responsive frontend dashboards with 1-click persona switching (Student, Mentor, HOD).\n"
            "• Created student activity submission portal supporting certificate uploads (PDF/PNG/SVG).\n"
            "• Developed Faculty Mentor review console to inspect pending department submissions.\n"
            "• Implemented degree milestone tiers (Bronze 50, Silver 100, Gold 150, Diamond 200)."
        ),
        "suggestions": (
            "• Incorporate automated OCR to assist faculty mentors in verifying student credentials.\n"
            "• Introduce discrepancy detection to flag mismatched student names or dates.\n"
            "• Provide real-time degree clearance progress meters on the student dashboard."
        ),
    },
    {
        "table": t2,
        "start_row": 1,
        "week_str": "Week 3\nfrom:\n14/09/2026\nto:\n19/09/2026",
        "tasks": (
            "• Integrated Google Gemini 2.5 Flash Vision for certificate OCR & metadata extraction.\n"
            "• Built discrepancy audit engine comparing certificate data against database profiles.\n"
            "• Constructed derived Institutional Knowledge Graph service mapping entities & relations.\n"
            "• Engineered deterministic fallback engine ensuring 100% offline platform availability.\n"
            "• Updated Mentor Review Console with side-by-side certificate proof & AI audit badges."
        ),
        "suggestions": (
            "• Ground all AI advisory recommendations in institutional graph evidence (GraphRAG).\n"
            "• Eliminate generative AI hallucinations regarding university policies or fake activities.\n"
            "• Design an official institutional credit transcript export format for degree clearance."
        ),
    },
    {
        "table": t2,
        "start_row": 5,
        "week_str": "Week 4\nfrom:\n21/09/2026\nto:\n26/09/2026",
        "tasks": (
            "• Deployed GraphRAG-powered AI Advisor providing grounded academic clearance roadmaps.\n"
            "• Built Student Visual Analytics suite: Domain Donut, Semester Velocity, & SGPA correlation.\n"
            "• Implemented official A4 Institutional Credit Transcript PDF engine with digital seal.\n"
            "• Completed end-to-end integration testing across Student, Mentor, and HOD workflows.\n"
            "• Finalized Phase I documentation, presentation deck, and prototype demonstration."
        ),
        "suggestions": (
            "• Simplify UI navigation for non-technical faculty evaluators using plain academic vocabulary.\n"
            "• Prepare Phase I demonstration showcasing certificate OCR, mentor review, & PDF transcript.\n"
            "• Verify edge cases for semester credit limits and category cap boundary conditions."
        ),
    }
]

def format_cell(cell, text, bold=False, font_size=8.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(font_size)
    run.font.bold = bold
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

# 5. Set CantSplit and balanced row heights
for t in [t1, t2]:
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    h0 = t.rows[0]._tr.get_or_add_trPr().find(qn('w:trHeight'))
    if h0 is not None:
        h0.set(qn('w:val'), '450')
        h0.set(qn('w:hRule'), 'atLeast')
    for r in t.rows[1:]:
        trPr = r._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
        h = trPr.find(qn('w:trHeight'))
        if h is None:
            h = OxmlElement('w:trHeight')
            trPr.append(h)
        h.set(qn('w:val'), '750')
        h.set(qn('w:hRule'), 'atLeast')

# 6. Populate table cells
for data in weeks_data:
    t = data["table"]
    start_row = data["start_row"]
    
    # Col 0: Week with date
    c0 = t.rows[start_row].cells[0]
    format_cell(c0, data["week_str"], bold=True, font_size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    
    # Col 1: Tasks achieved this week
    c1 = t.rows[start_row].cells[1]
    format_cell(c1, data["tasks"], bold=False, font_size=8, align=WD_ALIGN_PARAGRAPH.LEFT)
    
    # Col 2: Suggestions by guide
    c2 = t.rows[start_row].cells[2]
    format_cell(c2, data["suggestions"], bold=False, font_size=8, align=WD_ALIGN_PARAGRAPH.LEFT)
    
    # Rows 0..3 of each week: Name (Col 3), USN (Col 4), Signature (Col 5), Marks (Col 6)
    for r_offset in range(4):
        curr_row = start_row + r_offset
        # Name: Kept blank as requested
        c_name = t.rows[curr_row].cells[3]
        format_cell(c_name, "", font_size=8.5)
        # USN: Kept blank as requested
        c_usn = t.rows[curr_row].cells[4]
        format_cell(c_usn, "", font_size=8.5)
        # Student Signature: Kept blank for physical signatures
        c_sig = t.rows[curr_row].cells[5]
        format_cell(c_sig, "", font_size=8.5)
        # Marks (5): Kept completely blank as requested so faculty guide can award them
        c_marks = t.rows[curr_row].cells[6]
        format_cell(c_marks, "", font_size=8.5)

doc.save(TARGET_DOCX)
print(f"Successfully generated {TARGET_DOCX}")

# Generate PDF version
subprocess.run(['libreoffice', '--headless', '--convert-to', 'pdf', TARGET_DOCX, '--outdir', '.'], check=True)
print(f"Successfully generated {TARGET_PDF}")
